"""Generate the single STEP00 DOCX from its reviewed Markdown source.

Temporary planning-only script; remove it before merging planning into main.
"""
from __future__ import annotations
import re
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[2]
source = ROOT / "docs/v2_0_3_1_planning/00_ASTRA_V031_DELTA_AUDIT_RELEASE_READINESS_PLAN.md"
target = ROOT / "docs/v2_0_3_1_planning/docx/00_ASTRA_V031_DELTA_AUDIT_RELEASE_READINESS_PLAN.docx"
lines = source.read_text(encoding="utf-8").splitlines()
d = Document()
s = d.sections[0]
s.top_margin, s.bottom_margin = Inches(.76), Inches(.72)
s.left_margin, s.right_margin = Inches(.82), Inches(.80)
body = d.styles["Normal"]
body.font.name, body.font.size = "Aptos", Pt(9)
body.font.color.rgb = RGBColor(37, 53, 74)
body.paragraph_format.space_after = Pt(4)
for style, size in [("Title", 21), ("Heading 1", 14), ("Heading 2", 11), ("Heading 3", 10)]:
    st = d.styles[style]
    st.font.name = "Aptos"
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = RGBColor(13, 61, 127)
    st.paragraph_format.space_before = Pt(9)
    st.paragraph_format.space_after = Pt(4)
h = s.header.paragraphs[0]
h.text = "SOFTWARE FACTORY  |  ASTRA STEP00  |  V2-AI-VIDEO-COMPOSER"
h.runs[0].font.color.rgb = RGBColor(24, 87, 158)
h.runs[0].font.size = Pt(8)
f = s.footer.paragraphs[0]
f.text = "08 OKTOBER 2026  |  PLANNING ONLY  |  NO RELEASE"
f.alignment = WD_ALIGN_PARAGRAPH.RIGHT
f.runs[0].font.size = Pt(8)
f.runs[0].font.color.rgb = RGBColor(108, 119, 135)
i = 0
while i < len(lines):
    raw = lines[i].strip()
    if not raw:
        i += 1
        continue
    if raw.startswith("|") and raw.endswith("|"):
        block = []
        while i < len(lines) and lines[i].strip().startswith("|"):
            row = [x.strip().replace("\\|", "|") for x in lines[i].strip().strip("|").split("|")]
            if not row or all(re.fullmatch(r"[-: ]+", v or "-") for v in row):
                i += 1
                continue
            block.append(row)
            i += 1
        if not block:
            continue
        n = len(block[0])
        t = d.add_table(rows=1, cols=n)
        t.style = "Light Shading Accent 1"
        for j, x in enumerate(block[0]):
            t.rows[0].cells[j].text = x
        for row in block[1:]:
            cells = t.add_row().cells
            for j, x in enumerate(row[:n]):
                cells[j].text = x
        for k, row in enumerate(t.rows):
            pr = row._tr.get_or_add_trPr()
            pr.append(OxmlElement("w:cantSplit"))
            if k == 0:
                el = OxmlElement("w:tblHeader")
                el.set(qn("w:val"), "true")
                pr.append(el)
            for c in row.cells:
                for p in c.paragraphs:
                    p.paragraph_format.space_after = Pt(1)
                    for run in p.runs:
                        run.font.size = Pt(7.7)
                        if k == 0:
                            run.font.bold = True
        continue
    m = re.match(r"^(#{1,3})\\s+(.+)$", raw)
    if m:
        level = len(m.group(1))
        if level == 1:
            d.add_paragraph(m.group(2), style="Title")
        else:
            d.add_heading(m.group(2), level=level-1)
    elif raw.startswith("- "):
        d.add_paragraph(raw[2:], style="List Bullet")
    elif re.match(r"^\\d+\\.\\s+", raw):
        d.add_paragraph(re.sub(r"^\\d+\\.\\s+", "", raw), style="List Number")
    else:
        text = raw.replace("**", "").replace("`", "")
        d.add_paragraph(text)
    i += 1

target.parent.mkdir(parents=True, exist_ok=True)
d.save(target)
probe = Document(target)
assert len(probe.paragraphs) > 35 and len(probe.tables) >= 5
print(f"STEP00_DOCX_GENERATED file={target} bytes={target.stat().st_size} tables={len(probe.tables)}")
