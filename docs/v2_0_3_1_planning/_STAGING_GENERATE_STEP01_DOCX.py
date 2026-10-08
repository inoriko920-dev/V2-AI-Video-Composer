"""One-time STEP01 Markdown-to-DOCX converter. Remove before merge."""
from __future__ import annotations
import re
from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[2]
inp=ROOT/"docs/v2_0_3_1_planning/01_ASTRA_V031_IMPLEMENTATION_RELEASE_CONTRACT.md"
out=ROOT/"docs/v2_0_3_1_planning/docx/01_ASTRA_V031_IMPLEMENTATION_RELEASE_CONTRACT.docx"
lines=inp.read_text(encoding="utf-8").splitlines()
d=Document();s=d.sections[0]
s.page_height,s.page_width=Cm(29.7),Cm(21)
s.top_margin,s.bottom_margin=Cm(1.8),Cm(1.7)
s.left_margin,s.right_margin=Cm(1.8),Cm(1.65)
s.header_distance,s.footer_distance=Cm(.7),Cm(.6)
normal=d.styles["Normal"];normal.font.name="Aptos";normal.font.size=Pt(9)
normal.font.color.rgb=RGBColor(34,48,66)
normal.paragraph_format.space_after=Pt(5)
normal.paragraph_format.line_spacing=1.1
for name,size in [("Title",21),("Heading 1",14),("Heading 2",11),("Heading 3",10)]:
    st=d.styles[name];st.font.name="Aptos";st.font.size=Pt(size)
    st.font.bold=True;st.font.color.rgb=RGBColor(12,67,139)
    st.paragraph_format.space_before=Pt(10);st.paragraph_format.space_after=Pt(5)
p=s.header.paragraphs[0];p.text="V2-AI-VIDEO-COMPOSER  |  ASTRA STEP 01  |  08 OKTOBER 2026"
p.runs[0].font.size=Pt(7.7);p.runs[0].font.color.rgb=RGBColor(24,85,156)
p=s.footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
p.text="PLANNING ONLY - v0.3.1 kandidat - Halaman "
p.runs[0].font.size=Pt(7.7)
fld=OxmlElement("w:fldSimple");fld.set(qn("w:instr"),"PAGE");p._p.append(fld)
BT=chr(96)
def add(p,value):
    for token in re.split(r"(\*\*[^*]+\*\*|\x60[^\x60]+\x60)",value):
        if not token:continue
        code=token.startswith(BT) and token.endswith(BT)
        bold=token.startswith("**") and token.endswith("**")
        run=p.add_run(token[1:-1] if code else token[2:-2] if bold else token)
        if bold:run.bold=True
        if code:
            run.font.name="Consolas";run.font.size=Pt(7.7)
            run.font.color.rgb=RGBColor(18,79,144)
i=0;headings=0;tables=0
while i<len(lines):
    line=lines[i].strip()
    if not line or line=="---":i+=1;continue
    if line.startswith("|") and line.endswith("|"):
        rows=[]
        while i<len(lines) and lines[i].strip().startswith("|"):
            vals=[x.strip() for x in lines[i].strip().strip("|").split("|")]
            if not all(re.fullmatch(r"[-: ]+",v) for v in vals):rows.append(vals)
            i+=1
        if not rows:continue
        table=d.add_table(rows=0,cols=len(rows[0]));table.style="Light Shading Accent 1"
        table.alignment=WD_TABLE_ALIGNMENT.CENTER
        for vals in rows:
            cells=table.add_row().cells
            for j,item in enumerate(vals[:len(cells)]):
                add(cells[j].paragraphs[0],item)
        for k,row in enumerate(table.rows):
            trPr=row._tr.get_or_add_trPr();trPr.append(OxmlElement("w:cantSplit"))
            if k==0:
                h=OxmlElement("w:tblHeader");h.set(qn("w:val"),"true");trPr.append(h)
            for c in row.cells:
                for para in c.paragraphs:
                    para.paragraph_format.space_after=Pt(1)
                    para.paragraph_format.line_spacing=1.0
                    for run in para.runs:run.font.size=Pt(7.5)
        tables+=1
        continue
    m=re.match(r"^(#{1,3})\s+(.+)$",line)
    if m:
        level=len(m.group(1))
        p=d.add_paragraph(style="Title" if level==1 else f"Heading {level-1}")
        add(p,m.group(2));headings+=1;i+=1;continue
    if line.startswith("- "):
        p=d.add_paragraph(style="List Bullet");add(p,line[2:]);i+=1;continue
    if re.match(r"^\d+\.\s+",line):
        p=d.add_paragraph(style="List Number")
        add(p,re.sub(r"^\d+\.\s+","",line));i+=1;continue
    p=d.add_paragraph();add(p,line);i+=1
out.parent.mkdir(parents=True,exist_ok=True);d.save(out)
test=Document(out)
assert headings>=10,headings
assert len(test.tables)>=6,len(test.tables)
assert out.stat().st_size>30000,out.stat().st_size
print("STEP01_DOCX_STRUCTURAL_PASS",out.stat().st_size,"bytes",headings,"headings",len(test.tables),"tables",len(test.paragraphs),"paragraphs")
