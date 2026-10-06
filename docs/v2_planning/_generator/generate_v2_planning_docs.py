from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

HERE = Path(__file__).resolve().parent
OUT = HERE.parent
OUT.mkdir(parents=True, exist_ok=True)


def load_steps() -> list[tuple[str, str, dict]]:
    steps: list[tuple[str, str, dict]] = []
    for path in sorted(HERE.glob("steps_*.json")):
        raw = json.loads(path.read_text(encoding="utf-8"))
        for step, title, data in raw:
            steps.append((str(step), str(title), dict(data)))
    return steps


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(item)


def add_numbered(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.add_run(item)


def filename_for(step: str, title: str) -> str:
    safe = title.upper().replace(" ", "_").replace("/", "_").replace(",", "").replace("&", "AND")
    return f"STEP_{step}_V2_{safe}.docx"


def make_step_doc(step: str, title: str, data: dict) -> Path:
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.55)
    sec.bottom_margin = Inches(0.55)
    sec.left_margin = Inches(0.62)
    sec.right_margin = Inches(0.62)

    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(9.5)
    styles["Title"].font.name = "Aptos Display"
    styles["Title"].font.size = Pt(22)
    for style_name in ("Heading 1", "Heading 2"):
        styles[style_name].font.name = "Aptos Display"
    styles["Heading 1"].font.size = Pt(14)
    styles["Heading 2"].font.size = Pt(11.5)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(f"V2 AI VIDEO COMPOSER\nSTEP {step}")
    r.bold = True
    r.font.size = Pt(20)
    r.font.name = "Aptos Display"

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(title)
    r.bold = True
    r.font.size = Pt(14)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("PLANNING ONLY - NO APPLICATION CODE IMPLEMENTATION IN THIS STEP")
    r.bold = True
    r.font.size = Pt(9)

    table = doc.add_table(rows=4, cols=2)
    table.style = "Table Grid"
    meta = [
        ("Project", "V2-AI-Video-Composer"),
        ("Legacy source", "AI-Automatic-Video-Composer (read-only reference)"),
        ("Status", "PLANNING / source-of-truth"),
        ("Global rule", "All development happens only in V2 repository"),
    ]
    for i, (left, right) in enumerate(meta):
        table.cell(i, 0).text = left
        table.cell(i, 1).text = right
        set_cell_shading(table.cell(i, 0), "D9EAF7")
        for run in table.cell(i, 0).paragraphs[0].runs:
            run.bold = True

    sections = [
        ("1. Tujuan STEP", data["purpose"], "paragraph"),
        ("2. Fakta dan baseline yang digunakan", data["evidence"], "bullets"),
        ("3. Keputusan planning", data["decisions"], "bullets"),
        ("4. Pekerjaan rinci yang direncanakan", data["tasks"], "numbered"),
        ("5. Gate PASS / FAIL", data["gate"], "bullets"),
        ("6. Deliverable STEP", data["deliverables"], "bullets"),
    ]
    for heading, section_content, kind in sections:
        doc.add_heading(heading, level=1)
        if kind == "paragraph":
            doc.add_paragraph(section_content)
        elif kind == "bullets":
            add_bullets(doc, section_content)
        else:
            add_numbered(doc, section_content)

    doc.add_heading("7. Risiko utama dan mitigasi", level=1)
    for risk in data["risks"]:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(risk + ". ")
        p.add_run(
            "Mitigasi: wajib melalui adapter/gate, regression test, commit kecil, dan rollback point yang terdokumentasi."
        ).italic = True

    doc.add_heading("8. Aturan handoff ke AI berikutnya", level=1)
    add_bullets(
        doc,
        [
            f"Baca dokumen STEP {step} ini sebelum mengerjakan area terkait.",
            "Baca seluruh planning sebelumnya yang menjadi dependency keputusan STEP ini.",
            "Jangan menulis ke repo AI-Automatic-Video-Composer; repo tersebut hanya referensi read-only.",
            "Sebelum coding, pastikan CODING GATE pada STEP 13 telah dinyatakan PASS oleh pengguna/proses proyek.",
            "Setelah implementasi dimulai nanti, laporkan perubahan, bukti test, status gate, risiko tersisa, dan next exact action.",
        ],
    )
    doc.add_heading("9. Definition of Done untuk planning STEP ini", level=1)
    doc.add_paragraph(
        "Planning dianggap selesai ketika keputusan, ruang lingkup, gate, risiko, dan deliverable pada dokumen ini cukup spesifik sehingga AI lain dapat melanjutkan tanpa mengulang diskusi dasar atau menebak tujuan teknis."
    )

    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rr = footer.add_run(f"V2 AI Video Composer | STEP {step} | Planning Source of Truth")
    rr.font.size = Pt(8)

    if step == "02":
        sec.top_margin = Inches(0.50)
        sec.bottom_margin = Inches(0.50)
        styles["Normal"].font.size = Pt(9.25)
        for para in doc.paragraphs:
            pf = para.paragraph_format
            if para.style and para.style.name.startswith("Heading"):
                pf.space_before = Pt(6)
                pf.space_after = Pt(2)
            else:
                pf.space_after = Pt(0)
                pf.line_spacing = 1.0
        for tbl in doc.tables:
            for row in tbl.rows:
                for cell in row.cells:
                    for para in cell.paragraphs:
                        para.paragraph_format.space_after = Pt(0)

    path = OUT / filename_for(step, title)
    doc.save(path)
    return path


def make_index(steps: list[tuple[str, str, dict]]) -> Path:
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.6)
    sec.bottom_margin = Inches(0.6)
    sec.left_margin = Inches(0.65)
    sec.right_margin = Inches(0.65)
    doc.styles["Normal"].font.name = "Aptos"
    doc.styles["Normal"].font.size = Pt(9.5)

    doc.add_heading("V2 AI Video Composer - Master Planning Index", 0)
    doc.add_paragraph(
        "Dokumen ini adalah indeks source-of-truth untuk seluruh planning sebelum implementasi besar. Repo lama AI-Automatic-Video-Composer adalah read-only. Seluruh perubahan V2 hanya boleh terjadi di V2-AI-Video-Composer."
    )
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    for j, heading in enumerate(("STEP", "Dokumen", "Fokus", "Gate utama")):
        table.cell(0, j).text = heading
        set_cell_shading(table.cell(0, j), "D9EAF7")
        table.cell(0, j).paragraphs[0].runs[0].bold = True
    for step, title, data in steps:
        row = table.add_row().cells
        row[0].text = step
        row[1].text = title
        row[2].text = data["purpose"]
        row[3].text = data["gate"][0]

    for i, row in enumerate(table.rows):
        tr_pr = row._tr.get_or_add_trPr()
        tr_pr.append(OxmlElement("w:cantSplit"))
        if i == 0:
            header = OxmlElement("w:tblHeader")
            header.set(qn("w:val"), "true")
            tr_pr.append(header)
        for cell in row.cells:
            for para in cell.paragraphs:
                para.paragraph_format.space_after = Pt(0)
                for run in para.runs:
                    run.font.size = Pt(9.0)

    doc.add_heading("Aturan eksekusi", 1)
    add_bullets(
        doc,
        [
            "Planning STEP 00-13 harus dibaca sebagai satu rangkaian keputusan yang saling bergantung.",
            "Tidak ada refactor besar, dependency besar, migrasi schema, redesign UI, atau pergantian backend sebelum planning terkait disetujui dan CODING GATE STEP 13 PASS.",
            "Eksperimen mature backend dilakukan hanya di V2 dan harus removable/rollbackable.",
            "UI dan fitur saat ini adalah baseline kompatibilitas; mature repo digunakan untuk fondasi/engine/pola stabilitas, bukan untuk mengganti identitas aplikasi.",
            "Final release hanya boleh dipromosikan setelah QA Windows portable dan gate release STEP 11-13 PASS.",
        ],
    )
    path = OUT / "V2_MASTER_PLANNING_INDEX.docx"
    doc.save(path)
    return path


def main() -> None:
    steps = load_steps()
    if len(steps) != 14 or [step for step, _, _ in steps] != [f"{i:02d}" for i in range(14)]:
        raise SystemExit("Planning data incomplete: expected STEP 00-13 exactly once")
    created = [make_step_doc(step, title, data) for step, title, data in steps]
    created.append(make_index(steps))
    for path in created:
        print(path)


if __name__ == "__main__":
    main()
