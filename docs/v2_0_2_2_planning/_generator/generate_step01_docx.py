from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "01_REPOSITORY_HYGIENE_AND_WORKFLOW_SECURITY_PLAN.md"
OUTDIR = ROOT / "docx"
OUTPUT = OUTDIR / "01_V2_0.2.2_REPOSITORY_HYGIENE_AND_WORKFLOW_SECURITY_PLAN.docx"

doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.6)
sec.bottom_margin = Inches(0.6)
sec.left_margin = Inches(0.7)
sec.right_margin = Inches(0.7)
doc.styles["Normal"].font.name = "Aptos"
doc.styles["Normal"].font.size = Pt(10)

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = title.add_run("V2 AI Video Composer 0.2.2")
r.bold = True
r.font.size = Pt(24)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("STEP 01 — Repository Hygiene & Workflow Security Plan")
r.bold = True
r.font.size = Pt(18)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Planning-only source of truth — no application implementation")
r.italic = True

meta = doc.add_table(rows=5, cols=2)
meta.style = "Table Grid"
pairs = [
    ("Repository", "inoriko920-dev/V2-AI-Video-Composer"),
    ("Audited main", "bbc8135cdddff182e92d376c590914ccc125e3a4"),
    ("Frozen stable release", "v0.2.1 @ eb94efebf142ba8203dfc3ae3c5fa861222a0926"),
    ("Target line", "v0.2.2 maintenance / hardening"),
    ("STEP status", "STEP 01 — PASS / planning only"),
]
for row, pair in zip(meta.rows, pairs):
    row.cells[0].text, row.cells[1].text = pair
    trPr = row._tr.get_or_add_trPr()
    trPr.append(OxmlElement("w:cantSplit"))

def clean(text: str) -> str:
    text = text.replace("**", "").replace("`", "")
    return text.replace("$"+"{{ secrets.NAME }}", "GitHub Actions secrets expression")

pending = []
def flush():
    global pending
    if pending:
        doc.add_paragraph(" ".join(x.strip() for x in pending))
        pending = []

for raw in SOURCE.read_text(encoding="utf-8").splitlines():
    line = raw.rstrip()
    if not line:
        flush()
        continue
    if line.startswith("# "):
        continue
    if line.startswith("### "):
        flush()
        doc.add_heading(clean(line[4:]), level=2)
        continue
    if line.startswith("## "):
        flush()
        doc.add_heading(clean(line[3:]), level=1)
        continue
    if line.startswith("- "):
        flush()
        doc.add_paragraph(clean(line[2:]), style="List Bullet")
        continue
    if re.match(r"^\d+\.\s+", line):
        flush()
        doc.add_paragraph(clean(re.sub(r"^\d+\.\s+", "", line)), style="List Number")
        continue
    pending.append(clean(line))
flush()

for table in doc.tables:
    if table.rows:
        trPr = table.rows[0]._tr.get_or_add_trPr()
        header = OxmlElement("w:tblHeader")
        header.set(qn("w:val"), "true")
        trPr.append(header)

OUTDIR.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
