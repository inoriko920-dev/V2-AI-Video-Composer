from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "00_POST_RELEASE_AUDIT_AND_SCOPE.md"
OUTDIR = ROOT / "docx"
OUTPUT = OUTDIR / "00_V2_0.2.2_POST_RELEASE_AUDIT_AND_SCOPE.docx"

def add_cant_split(row):
    trPr = row._tr.get_or_add_trPr()
    if trPr.find(qn("w:cantSplit")) is None:
        trPr.append(OxmlElement("w:cantSplit"))

doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.65)
sec.bottom_margin = Inches(0.65)
sec.left_margin = Inches(0.75)
sec.right_margin = Inches(0.75)
doc.styles["Normal"].font.name = "Aptos"
doc.styles["Normal"].font.size = Pt(10)

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run("V2 AI Video Composer 0.2.2")
run.bold = True
run.font.size = Pt(25)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("STEP 00 — Post-Release Audit & Scope")
r.bold = True
r.font.size = Pt(20)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Maintenance / Hardening Planning — Source of Truth")
r.italic = True

meta = doc.add_table(rows=6, cols=2)
meta.style = "Table Grid"
meta_data = [
    ("Repository", "inoriko920-dev/V2-AI-Video-Composer"),
    ("Baseline main", "e9ff8540f3cad74035a592f4d4670c136ea64652"),
    ("Frozen stable release", "v0.2.1 @ eb94efebf142ba8203dfc3ae3c5fa861222a0926"),
    ("Previous stable", "v0.2.0 @ c6ad75308d302cd816301a2a7a34ebe7eb4148a7"),
    ("Target planning line", "v0.2.2"),
    ("STEP status", "STEP 00 — PASS / Planning only"),
]
for row, pair in zip(meta.rows, meta_data):
    row.cells[0].text, row.cells[1].text = pair
    add_cant_split(row)

lines = SOURCE.read_text(encoding="utf-8").splitlines()
paragraph = []

def flush():
    global paragraph
    if paragraph:
        doc.add_paragraph(" ".join(x.strip() for x in paragraph).strip())
        paragraph = []

for raw in lines:
    line = raw.rstrip()
    if not line:
        flush()
        continue
    if line.startswith("# "):
        continue
    if line.startswith("### "):
        flush(); doc.add_heading(line[4:], level=3); continue
    if line.startswith("## "):
        flush(); doc.add_heading(line[3:], level=1); continue
    if re.match(r"^\d+\.\s+", line):
        flush(); doc.add_paragraph(re.sub(r"^\d+\.\s+", "", line), style="List Number"); continue
    if line.startswith("- "):
        flush(); doc.add_paragraph(line[2:], style="List Bullet"); continue
    if line.startswith("**") and line.endswith("**"):
        flush()
        p = doc.add_paragraph()
        r = p.add_run(line.strip("*"))
        r.bold = True
        continue
    paragraph.append(line)
flush()

for table in doc.tables:
    for row in table.rows:
        add_cant_split(row)

OUTDIR.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
