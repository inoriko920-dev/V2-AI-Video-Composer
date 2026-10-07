from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "04_WINDOWS_PACKAGING_DEPENDENCY_MAINTENANCE_PLAN.md"
OUTDIR = ROOT / "docx"
OUTPUT = OUTDIR / "04_V2_0.2.2_WINDOWS_PACKAGING_DEPENDENCY_MAINTENANCE_PLAN.docx"

doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.55)
sec.bottom_margin = Inches(0.55)
sec.left_margin = Inches(0.62)
sec.right_margin = Inches(0.62)
doc.styles["Normal"].font.name = "Aptos"
doc.styles["Normal"].font.size = Pt(9.2)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("V2 AI Video Composer 0.2.2")
r.bold = True
r.font.size = Pt(23)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("STEP 04 — Windows Packaging / Dependency Maintenance Plan")
r.bold = True
r.font.size = Pt(18)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Planning-only source of truth — no application implementation")
r.italic = True

lines = SOURCE.read_text(encoding="utf-8").splitlines()
paragraph = []
table_rows = []

def clean(text: str) -> str:
    return text.replace(chr(96), "").replace("**", "")

def flush() -> None:
    global paragraph
    if paragraph:
        doc.add_paragraph(clean(" ".join(x.strip() for x in paragraph).strip()))
        paragraph = []

def flush_table() -> None:
    global table_rows
    if not table_rows:
        return
    rows = [r for r in table_rows if not all(set(c.strip()) <= {"-", ":"} for c in r)]
    if rows:
        table = doc.add_table(rows=1, cols=len(rows[0]))
        table.style = "Table Grid"
        for i, v in enumerate(rows[0]):
            table.rows[0].cells[i].text = clean(v.strip())
        for row in rows[1:]:
            cells = table.add_row().cells
            for i, v in enumerate(row):
                cells[i].text = clean(v.strip())
    table_rows = []

for raw in lines:
    line = raw.rstrip()
    if line.startswith("|") and line.endswith("|"):
        flush()
        table_rows.append([x for x in line.strip("|").split("|")])
        continue
    flush_table()
    if not line:
        flush()
        continue
    if line.startswith("# "):
        continue
    if line.startswith("### "):
        flush(); doc.add_heading(clean(line[4:]), level=2); continue
    if line.startswith("## "):
        flush(); doc.add_heading(clean(line[3:]), level=1); continue
    if re.match(r"^\d+\.\s+", line):
        flush(); doc.add_paragraph(clean(re.sub(r"^\d+\.\s+", "", line)), style="List Number"); continue
    if line.startswith("- "):
        flush(); doc.add_paragraph(clean(line[2:]), style="List Bullet"); continue
    if line.startswith("**") and line.endswith("**"):
        flush()
        p = doc.add_paragraph()
        r = p.add_run(clean(line))
        r.bold = True
        continue
    paragraph.append(line)
flush()
flush_table()

for table in doc.tables:
    if table.rows:
        trPr = table.rows[0]._tr.get_or_add_trPr()
        header = OxmlElement("w:tblHeader")
        header.set(qn("w:val"), "true")
        trPr.append(header)
    for row in table.rows:
        trPr = row._tr.get_or_add_trPr()
        trPr.append(OxmlElement("w:cantSplit"))

OUTDIR.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
