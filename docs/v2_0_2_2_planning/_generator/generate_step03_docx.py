from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "03_SCHEMA_V4_PERSISTENCE_RECOVERY_STRESS_PLAN.md"
OUTDIR = ROOT / "docx"
OUTPUT = OUTDIR / "03_V2_0.2.2_SCHEMA_V4_PERSISTENCE_RECOVERY_STRESS_PLAN.docx"

doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.55)
sec.bottom_margin = Inches(0.55)
sec.left_margin = Inches(0.65)
sec.right_margin = Inches(0.65)
doc.styles["Normal"].font.name = "Aptos"
doc.styles["Normal"].font.size = Pt(9.5)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("V2 AI Video Composer 0.2.2")
r.bold = True
r.font.size = Pt(24)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("STEP 03 — Schema-v4 Persistence / Recovery Stress Plan")
r.bold = True
r.font.size = Pt(19)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Planning-only source of truth — no application implementation")
r.italic = True

lines = SOURCE.read_text(encoding="utf-8").splitlines()
paragraph = []

def clean(text: str) -> str:
    return text.replace(chr(96), "").replace("**", "")

def flush() -> None:
    global paragraph
    if paragraph:
        doc.add_paragraph(clean(" ".join(x.strip() for x in paragraph).strip()))
        paragraph = []

for raw in lines:
    line = raw.rstrip()
    if not line:
        flush()
        continue
    if line.startswith("# "):
        continue
    if line.startswith("### "):
        flush(); doc.add_heading(clean(line[4:]), level=2); continue
    if line.startswith("## "):
        flush(); doc.add_heading(clean(line[3:]), level=1); continue
    if re.match(r"^\d+\.\s+", line):
        flush(); doc.add_paragraph(clean(re.sub(r"^\d+\.\s+", "", line)), style="List Number"); continue
    if line.startswith("- "):
        flush(); doc.add_paragraph(clean(line[2:]), style="List Bullet"); continue
    if line.startswith("**") and line.endswith("**"):
        flush()
        p = doc.add_paragraph()
        r = p.add_run(clean(line))
        r.bold = True
        continue
    paragraph.append(line)
flush()

for table in doc.tables:
    if table.rows:
        trPr = table.rows[0]._tr.get_or_add_trPr()
        header = OxmlElement("w:tblHeader")
        header.set(qn("w:val"), "true")
        trPr.append(header)
    for row in table.rows:
        trPr = row._tr.get_or_add_trPr()
        trPr.append(OxmlElement("w:cantSplit"))

OUTDIR.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
