from pathlib import Path
import hashlib, json, zipfile
from datetime import datetime
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_ORIENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_BREAK
from PIL import Image

ROOT=Path('/mnt/data')
PNG=ROOT/'V2_AI_Video_Composer_STEP03_R_UI_DRAFT_REVIEW'
DOCX=ROOT/'03_V2_0.3.0_FINAL_APPROVED_UI_DIALOG_REFERENCES.docx'
SOURCES=[PNG/'UI-REC-DLG-01.png',PNG/'UI-REC-DLG-02.png',PNG/'UI-REC-DLG-03.png']
ORIGINAL=PNG/'UI-013_ACTUAL_REFERENCE.png'
ALL=SOURCES+[ORIGINAL]
for path in ALL:
    assert path.exists(), path
    with Image.open(path) as im:
        assert im.size == (1920,1080), (path,im.size)
    assert path.stat().st_size>100_000
hashes={p.name:hashlib.sha256(p.read_bytes()).hexdigest() for p in ALL}

D=Document()
s=D.sections[0]; s.orientation=WD_ORIENT.LANDSCAPE
s.page_width=Inches(11.69); s.page_height=Inches(8.27)
s.left_margin=Inches(.73);s.right_margin=Inches(.73);s.top_margin=Inches(.73);s.bottom_margin=Inches(.58)
s.header_distance=Inches(.28);s.footer_distance=Inches(.27)
styles=D.styles
for st in ['Normal','Title','Heading 1','Heading 2']:
    sty=styles[st];sty.font.name='Aptos';sty.font.color.rgb=RGBColor(26,39,58)
styles['Normal'].font.size=Pt(9)
styles['Normal'].paragraph_format.space_after=Pt(5)
for name,pt in [('Title',24),('Heading 1',15),('Heading 2',11)]:
    styles[name].font.size=Pt(pt); styles[name].font.bold=True
styles['Heading 1'].font.color.rgb=RGBColor(31,83,172)

hdr=s.header.paragraphs[0];hdr.text='V2-AI-VIDEO-COMPOSER  /  STEP 03-R  /  APPROVED UI REFERENCE'
hdr.style='Caption';hdr.alignment=WD_ALIGN_PARAGRAPH.RIGHT
hdr.runs[0].font.size=Pt(7);hdr.runs[0].font.color.rgb=RGBColor(95,111,134)
foot=s.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.RIGHT
foot.add_run('Referensi visual, bukan bukti fitur telah diimplementasikan    |    Halaman ')
field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE');foot._p.append(field)
for run in foot.runs:run.font.size=Pt(7)

def p(text='',bold=False,style=None):
    q=D.add_paragraph(style=style)
    q.add_run(text).bold=bold
    return q

def kv(key,value):
    q=D.add_paragraph(style='Normal')
    a=q.add_run(key+'  '); a.bold=True; a.font.color.rgb=RGBColor(30,79,158)
    q.add_run(value)
    return q

def title(x):D.add_heading(x,level=1)
def subtitle(x):D.add_heading(x,level=2)

def bullets(xs):
    for x in xs:D.add_paragraph(x,style='List Bullet')

def table(rows, widths=None):
    t=D.add_table(rows=1, cols=len(rows[0]));t.style='Light Shading Accent 1';t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for j,key in enumerate(rows[0]):t.rows[0].cells[j].text=str(key)
    for r in rows[1:]:
        cells=t.add_row().cells
        for j,c in enumerate(r):cells[j].text=str(c)
    for row in t.rows:
        for cell in row.cells:
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for par in cell.paragraphs:
                for run in par.runs:run.font.size=Pt(8)
    return t

def imagepage(path, label, desc, follow):
    D.add_page_break();title(label)
    q=D.add_paragraph();q.alignment=WD_ALIGN_PARAGRAPH.CENTER
    q.add_run().add_picture(str(path),width=Inches(9.25))
    q=D.add_paragraph(desc);q.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for run in q.runs:run.italic=True;run.font.size=Pt(8)
    for item in follow:
        q=D.add_paragraph('• '+item)
        q.paragraph_format.space_after=Pt(2)
        for r in q.runs:r.font.size=Pt(8)

D.add_heading('Referensi UI Final yang Disetujui',0)
p('STEP 03-R — 3 Dialog Pemulihan Proyek | V2-AI-Video-Composer',bold=True)
p('Dokumen otoritatif untuk implementasi visual setelah persetujuan eksplisit pengguna. Hanya referensi; tidak mengesahkan perubahan kode, implementasi autosave/recovery, ataupun STEP 04–06.')
D.add_paragraph('')
table([
    ['Identitas','Keputusan'],
    ['Repositori','inoriko920-dev/V2-AI-Video-Composer (khusus V2)'],
    ['Baseline STEP03-R','PR #53 merged main 1354d386bcf9f47c7974309c4a8a1bf98584bcb1'],
    ['Versi stabil dibekukan','v0.2.2 @ eeabf1ffbfac6cbb6cafc82116fc3599bdbc2bef'],
    ['Cakupan final','3 dialog recovery saja; editor/home/UI lama tetap 1:1'],
    ['Status persetujuan','Pengguna menyetujui 3 PNG secara eksplisit dalam percakapan 8 Oktober 2026'],
    ['Otorisasi','Finalisasi referensi dan unggah ke repo V2; dilarang mulai coding']
])
D.add_paragraph('')
subtitle('Kriteria keberterimaan desain')
bullets([
  'Tiga gambar individual 1920 × 1080; bukan kolase atau pengganti seluruh editor.',
  'Semua modal berada di atas screenshot editor aktual UI-013; toolbar, pratinjau, panel aset/properti/AI, dan timeline utama tidak dirancang ulang.',
  'Seluruh copy utama, tombol dan arah tindakan berbahasa Indonesia; Batal adalah pilihan aman/default.',
  'Tidak ada Restore otomatis, penghapusan cadangan tanpa izin, klaim data sudah tersimpan yang keliru, timestamp fiktif, atau path/API key terbuka.',
  'Indikator autosave dan dialog perubahan belum disimpan menggunakan komponen lama, tanpa desain jendela baru.'
])
subtitle('Batas hak akses dan tahap')
p('Dokumen ini melengkapi gate gambar STEP 03-R. STEP 04 baru dapat direncanakan atas perintah pengguna tersendiri setelah berkas benar-benar disimpan dan diverifikasi di GitHub. STEP 06 implementasi SOL tetap membutuhkan seluruh planning DOCX dan otorisasi independen.')

imagepage(ORIGINAL,'Acuan UI-013: screenshot editor asli',
'Sumber: GitHub Actions CI STEP09, workflow run 37734080472, artifact step09-ui-actual (11531031921).',[
 'UI-002_ACTUAL adalah Home, bukan editor; UI-013_ACTUAL adalah acuan tata letak yang benar.',
 'Tulisan autosave pada screenshot contoh tidak membuktikan autosave produksi telah diimplementasikan.'
])

imagepage(SOURCES[0],'UI-REC-DLG-01 — Cadangan Proyek Ditemukan',
'Gambar 1: pemulihan kandidat valid (RECOVERABLE_VERIFIED).',[
 'Tiga tombol: Pulihkan Cadangan; Gunakan Proyek Tersimpan; Batal (fokus/default).',
 'Restore harus cek ulang snapshot dan membuat backup unik sebelum penulisan disk; Batal tidak mengubah file/sesi.',
 'Tidak boleh menganggap cadangan otomatis sebagai Save manual atau menghapus snapshot lain.'
])

imagepage(SOURCES[1],'UI-REC-DLG-02 — Cadangan Perlu Diperiksa',
'Gambar 2: baseline berubah / provenance lama / kandidat tidak pasti (RECOVERABLE_UNCERTAIN).',[
 'Tiga tombol: Periksa dan Pulihkan; Gunakan Proyek Tersimpan; Batal (fokus/default).',
 'Setelah Periksa dan Pulihkan, wajib konfirmasi kedua PADA KOMPONEN YANG SAMA: Ya, Pulihkan Cadangan / Batal.',
 'Periksa ulang hash ketika tindakan; bila berubah, tampilkan Data proyek berubah. Periksa ulang. tanpa melakukan restore.'
])

imagepage(SOURCES[2],'UI-REC-DLG-03 — Cadangan Tidak Dapat Dibaca',
'Gambar 3: cadangan tidak valid/tidak terbaca, memakai pola QMessageBox native.',[
 'Dua tombol: Buka Proyek Tersimpan; Batal (fokus/default). Tidak boleh ada tombol Pulihkan.',
 'Snapshot invalid tetap tersedia untuk pemeriksaan; jangan ada shortcut Hapus Cadangan.',
 'Pada partial commit, pesan harus menjelaskan keadaan disk aktual, tidak selalu mengklaim disk tidak berubah.'
])

D.add_page_break();title('Kontrak perilaku dan cakupan yang dibekukan')
subtitle('Keamanan recovery dan guard')
table([
 ['Keadaan','Prinsip yang wajib dipertahankan'],
 ['Kandidat valid','Pulihkan hanya setelah revalidasi disk & cadangan; buat backup proyek sebelum mengganti.'],
 ['Konflik','Konfirmasi tambahan; snapshot hash/ownership dicek lagi di saat aksi.'],
 ['Cadangan rusak','Blokir restore; file cadangan tidak otomatis dibuang.'],
 ['Batal / Esc','Tidak mengubah disk, proyek aktif, snapshot, atau dirty state.'],
 ['Save/Discard/Cancel','Gunakan GuardedMainWindow QMessageBox yang sudah ada, terlebih dahulu jika diperlukan.'],
 ['Snapshot selesai','Tetap dirty dan tetap bisa Undo/Redo tanpa status Saved palsu.'],
 ['Restore parsial','Jangan menjanjikan disk asli tidak berubah jika replace durable sempat terjadi.']
])
subtitle('Indikator pada QStatusBar lama, bukan desain dialog baru')
table([
 ['Kondisi','Copy yang wajib digunakan'],
 ['Menunggu jeda','Cadangan otomatis menunggu jeda.'],
 ['Sedang menulis','Membuat cadangan otomatis…'],
 ['Selesai','Cadangan otomatis dibuat.'],
 ['Gagal','Cadangan otomatis gagal — simpan manual tersedia.'],
 ['Proyek belum punya path','Simpan proyek untuk mengaktifkan cadangan otomatis.'],
 ['Save manual berhasil','Proyek disimpan.']
])
subtitle('Cakupan yang dilarang berubah pada tahap ini')
p('Tidak mengubah 42 referensi UI lama; halaman Home, toolbar, panel aset, preview, properti/AI, timeline, shortcut; kode Python/PySide6, skema proyek, scheduler, FFmpeg, dependensi, test, build atau rilis; repo versi asli dan tag stabil v0.2.2.')

D.add_page_break();title('Bukti persetujuan, integritas aset, dan handoff')
subtitle('Persetujuan manusia')
p('Pernyataan pengguna dalam percakapan (8 Oktober 2026): “Saya secara eksplisit menyetujui ketiga gambar UI STEP 03-R V2-AI-Video-Composer sebagai desain final. Buat DOCX referensi UI final berisi gambar, masukkan dokumen dan ketiga PNG ke GitHub V2, lalu lakukan pemeriksaan. Jangan mulai coding.”')
p('Persetujuan berlaku untuk ketiga visual dalam dokumen ini, bukan untuk mulai coding atau mengubah UI editor utama. Konfirmasi kedua pada dialog konflik tetap menjadi spesifikasi perilaku, bukan gambar keempat.')
subtitle('Manifest aset yang disetujui (SHA-256)')
for filename,digest in hashes.items():
    kv(filename, digest)
subtitle('Lokasi arsip dalam repository')
bullets([
  'docs/v2_0_3_0_planning/ui_step03_approved/UI-REC-DLG-01.png',
  'docs/v2_0_3_0_planning/ui_step03_approved/UI-REC-DLG-02.png',
  'docs/v2_0_3_0_planning/ui_step03_approved/UI-REC-DLG-03.png',
  'docs/v2_0_3_0_planning/ui_step03_approved/UI-013_ACTUAL_REFERENCE.png (acuan komparasi)',
  'docs/v2_0_3_0_planning/docx/03_V2_0.3.0_FINAL_APPROVED_UI_DIALOG_REFERENCES.docx',
  'docs/v2_0_3_0_planning/03_STEP03_APPROVAL_RECORD.md (bukti gate dan catatan audit)'
])
subtitle('Syarat sebelum status gate diubah menjadi PASS')
p('Commit seluruh file di cabang khusus, verifikasi SHA PNG dan DOCX yang terbaca, pastikan diff hanya dokumen/aset, tunggu CI, backend, CodeQL dan Windows acceptance sukses, baru gabung PR ke main. Bila tidak berhasil, laporkan PARTIAL/HOLD secara transparan. Setelah gate PASS, tetap HARUS BERHENTI menunggu perintah berikutnya untuk STEP 04.')

# explicit embedded count after save
D.core_properties.title='V2-AI-Video-Composer | STEP03-R Final Approved UI Dialog References'
D.core_properties.subject='Three approved UI dialogs; original editor preserved'
D.core_properties.keywords='STEP03-R, recovery, V2-AI-Video-Composer, approved UI'
D.save(DOCX)
with zipfile.ZipFile(DOCX) as z:
   emb=[a for a in z.namelist() if a.startswith('word/media/')]
   assert len(emb)>=4,(len(emb),emb)
   assert z.testzip() is None
print('DOCX_PATH',DOCX)
print('DOCX_BYTES',DOCX.stat().st_size)
print('DOCX_SHA256',hashlib.sha256(DOCX.read_bytes()).hexdigest())
print('EMBEDDED_MEDIA_COUNT',len(emb))
for k,v in hashes.items():print('IMAGE_SHA256',k,v)
manifest={'sha256':hashes,'docx_sha256':hashlib.sha256(DOCX.read_bytes()).hexdigest(),'images_approved':True,'approval_date_wib':'2026-10-08','repository':'inoriko920-dev/V2-AI-Video-Composer','expected_files':list(hashes)}
(ROOT/'STEP03_R_APPROVED_FILE_MANIFEST.json').write_text(json.dumps(manifest,indent=2,ensure_ascii=False),encoding='utf-8')