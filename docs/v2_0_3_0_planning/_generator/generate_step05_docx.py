#!/usr/bin/env python3
"""Create STEP05 planning DOCX from repository Markdown, no application source changes."""
from __future__ import annotations
import argparse
from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

DEFAULT_SRC=Path('docs/v2_0_3_0_planning/05_IMPLEMENTATION_READINESS_AND_SOL_WAVE_AUTHORIZATION_PLAN.md')
DEFAULT_DST=Path('docs/v2_0_3_0_planning/docx/05_V2_0.3.0_IMPLEMENTATION_READINESS_AND_SOL_WAVE_PLAN.docx')

def color(cell,fill):
    tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),fill); tcPr.append(shd)
def prevent_split(row):
    trPr=row._tr.get_or_add_trPr(); el=OxmlElement('w:cantSplit');trPr.append(el)
def row_header(row):
    trPr=row._tr.get_or_add_trPr();el=OxmlElement('w:tblHeader');el.set(qn('w:val'),'true');trPr.append(el)

def write_inlines(p,txt):
    # Plain text source, with bold and inline-code rendered natively.
    s=re.split(r'(\*\*.*?\*\*|`[^`]+`)',txt)
    for part in s:
        if not part:continue
        if part.startswith('**') and part.endswith('**'):
            p.add_run(part[2:-2]).bold=True
        elif part.startswith('`') and part.endswith('`'):
            r=p.add_run(part[1:-1]);r.font.name='Consolas';r.font.size=Pt(8)
        else:
            p.add_run(part)

def main():
    ap=argparse.ArgumentParser();ap.add_argument('--source',type=Path,default=DEFAULT_SRC);ap.add_argument('--output',type=Path,default=DEFAULT_DST)
    args=ap.parse_args(); raw=args.source.read_text(encoding='utf-8');assert len(raw)>14000
    doc=Document();section=doc.sections[0];section.page_height=Inches(11.69);section.page_width=Inches(8.27)
    section.top_margin=Inches(.68);section.bottom_margin=Inches(.64);section.left_margin=Inches(.7);section.right_margin=Inches(.7)
    section.header_distance=Inches(.32);section.footer_distance=Inches(.34)
    styles=doc.styles; normal=styles['Normal'];normal.font.name='Aptos';normal.font.size=Pt(8.8);normal.font.color.rgb=RGBColor(32,49,67)
    normal.paragraph_format.space_after=Pt(4);normal.paragraph_format.line_spacing=1.08
    for name,size,space in [('Title',23,9),('Heading 1',12.5,6),('Heading 2',10.5,5),('Heading 3',9.5,4)]:
        st=styles[name];st.font.name='Aptos Display';st.font.size=Pt(size);st.font.bold=True;st.font.color.rgb=RGBColor(18,77,140)
        st.paragraph_format.space_before=Pt(space+4);st.paragraph_format.space_after=Pt(space);st.paragraph_format.keep_with_next=True
    h=section.header.paragraphs[0];h.alignment=WD_ALIGN_PARAGRAPH.RIGHT;rr=h.add_run('V2 AI VIDEO COMPOSER   /   STEP 05');rr.font.size=Pt(8);rr.font.bold=True;rr.font.color.rgb=RGBColor(36,99,157)
    f=section.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    fr=f.add_run('PLANNING ONLY  •  08 OKT 2026  |  HALAMAN ');fr.font.size=Pt(7);fr.font.color.rgb=RGBColor(91,104,119)
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');f._p.append(field)
    lines=raw.splitlines();i=0;table_count=0
    while i<len(lines):
        ln=lines[i].strip()
        if not ln: i+=1;continue
        if ln.startswith('|'):
            rawrows=[]
            while i<len(lines) and lines[i].strip().startswith('|'):
                rawrows.append([v.strip() for v in lines[i].strip().strip('|').split('|')]);i+=1
            if len(rawrows)<2:continue
            rows=[r for r in rawrows if not all(re.fullmatch(r'[:\- ]+',v or '') for v in r)]
            cols=max(map(len,rows));table=doc.add_table(rows=0,cols=cols);table.style='Table Grid';table.alignment=WD_TABLE_ALIGNMENT.CENTER;table.autofit=True
            for ix,row in enumerate(rows):
                rc=table.add_row();prevent_split(rc)
                if ix==0:row_header(rc)
                for j,c in enumerate(rc.cells):
                    c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    if ix==0: color(c,'E9F2FA')
                    elif ix%2==0:color(c,'F7FAFD')
                    pp=c.paragraphs[0];pp.paragraph_format.space_after=Pt(2);pp.paragraph_format.space_before=Pt(2)
                    write_inlines(pp,row[j] if j<len(row) else '')
                    for r in pp.runs:r.font.size=Pt(7.6 if cols>=3 else 8)
                    if ix==0:
                        for r in pp.runs:r.bold=True
            table_count+=1;doc.add_paragraph('');continue
        if ln.startswith('# '):
            p=doc.add_paragraph(style='Title');write_inlines(p,ln[2:]);i+=1;continue
        if ln.startswith('## '):
            p=doc.add_paragraph(style='Heading 1');write_inlines(p,ln[3:]);i+=1;continue
        if ln.startswith('### '):
            p=doc.add_paragraph(style='Heading 2');write_inlines(p,ln[4:]);i+=1;continue
        if ln.startswith('- '):
            p=doc.add_paragraph(style='List Bullet');write_inlines(p,ln[2:]);i+=1;continue
        if re.match(r'^\d+\. ',ln):
            p=doc.add_paragraph(style='List Number');write_inlines(p,re.sub(r'^\d+\. ','',ln));i+=1;continue
        p=doc.add_paragraph();write_inlines(p,ln);i+=1
    props=doc.core_properties;props.title='V2-AI-Video-Composer — STEP05 Implementation Readiness & SOL Waves';props.subject='Planning-only Definition of Ready; recovery v0.3.0';props.keywords='STEP05,ASTRA,SOL,Recovery,Windows,No coding'
    args.output.parent.mkdir(parents=True,exist_ok=True);doc.save(args.output)
    assert table_count>=5,table_count
    from zipfile import ZipFile
    with ZipFile(args.output) as z:assert z.testzip() is None
    print(f'Generated {args.output} {args.output.stat().st_size} bytes, tables {table_count}, paragraphs {len(doc.paragraphs)}')
if __name__=='__main__':main()